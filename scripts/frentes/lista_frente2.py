"""Lista da Frente 2 (mobilizar): cidades pró-Lula com 10 a 50 mil habitantes.

Entra: município com 10.000 a 49.999 habitantes no Censo 2022 (IBGE), onde Lula teve mais
de 50% dos votos entre ele e Flávio no 1º turno de 2026 e venceu o 2º turno de 2022.
Ordem: saldo potencial = ausentes × vantagem de Lula (fórmula na planilha).

Entradas: docs/data/municipios.json, site/data/municipios.csv e site/data/regioes.json
(rode antes scripts/site/build_dados.py).
Saída: outputs/frente2_cidades_10a50mil_habitantes.xlsx (com fórmulas), copiada para
site/data/ para download no site.
Depois de gerar, recalcule as fórmulas em Excel ou LibreOffice.

Uso, na raiz do repositório: python3 scripts/frentes/lista_frente2.py   (requer openpyxl)
"""
import csv
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

REPO = Path(__file__).resolve().parents[2]
OUT = REPO / "outputs/frente2_cidades_10a50mil_habitantes.xlsx"
D = json.load(open(REPO / 'docs/data/municipios.json', encoding='utf-8'))
C = {r['ibge']: r for r in csv.DictReader(open(REPO / 'site/data/municipios.csv', encoding='utf-8'))}
REG = json.load(open(REPO / 'site/data/regioes.json', encoding='utf-8'))['imediatas']
tidy = lambda s: ' '.join(w.lower() if w in ('De','Da','Do','Das','Dos','E') else w for w in s.split(' '))

rows = []
for k, d in D.items():
    pop = int(C[k]['pop_2022']) if C.get(k, {}).get('pop_2022') else 0
    if d.get('l22t2') is None or not (10000 <= pop < 50000): continue
    L = d['vl'] / (d['vl'] + d['vf'])
    if not (L > 0.5 and d['l22t2'] > 50): continue
    gov = ''
    if d.get('g2t'):
        gov = 'sim' if d.get('g2ts') == 'confirmado' else 'incerto (TSE pode anular votos de candidato)'
    rows.append(dict(ibge=k, n=tidy(d['n']), uf=d['uf'], rg=d['rg'], arr=REG.get(C[k]['regiao_imediata'], [''])[0], pop=pop, apt=d['apt'], aus=d['aus'],
                     a22=d['a22'] / 100, dab=d.get('dab22'), vl=d['vl'], vf=d['vf'], l2t=d['l22t2'] / 100,
                     i70=d.get('i70'), pp=d.get('pp') or '', pb=d.get('pb') or '', gov=gov or 'não',
                     saldo=d['aus'] * (2 * L - 1)))
rows.sort(key=lambda r: -r['saldo'])
N = len(rows); last = N + 1

wb = Workbook()
F = 'Arial'
hdr_font = Font(name=F, bold=True, size=10); body = Font(name=F, size=10); blue = Font(name=F, size=10, color='0000FF')
hdr_fill = PatternFill('solid', fgColor='E8EAE4'); yellow = PatternFill('solid', fgColor='FFFF00')
thin = Border(bottom=Side(style='thin', color='BBBBBB'))

# ---------------- Cidades ----------------
ws = wb.active; ws.title = 'Cidades'
cols = [
 ('Posição', 8, '0'), ('Prioridade', 10, '@'), ('Município', 26, '@'), ('UF', 5, '@'), ('Região', 12, '@'),
 ('Região imediata (IBGE 2017)', 26, '@'),
 ('Saldo potencial para Lula (votos)', 14, '#,##0'), ('Saldo com +1 p.p. de comparecimento (votos)', 15, '#,##0'),
 ('Eleitores aptos 2026', 12, '#,##0'), ('Ausentes 1º turno 2026', 12, '#,##0'), ('Abstenção 1º turno 2026', 11, '0.0%'),
 ('Abstenção 1º turno 2022', 11, '0.0%'), ('Abstenção subiu desde 2022?', 11, '@'),
 ('Variação da abstenção do 1º para o 2º turno de 2022 (p.p.)', 15, '+0.0;-0.0;0.0'),
 ('Votos Lula 1º turno 2026', 12, '#,##0'), ('Votos Flávio 1º turno 2026', 12, '#,##0'), ('Lula entre os dois 2026', 11, '0.0%'),
 ('Lula no 2º turno 2022', 11, '0.0%'), ('Vantagem de Lula (Lula − Flávio)', 12, '+0.0%;-0.0%'), ('Eleitores com 70 anos ou mais', 11, '0.0%'),
 ('Partido do prefeito (2024)', 13, '@'), ('Campo do prefeito (2024)', 11, '@'), ('2º turno para governador no estado', 14, '@'), ('Código IBGE', 10, '@'),
 ('População (Censo 2022)', 12, '#,##0'),
]
for j, (h, w, _) in enumerate(cols, 1):
    c = ws.cell(row=1, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill
    c.alignment = Alignment(wrap_text=True, vertical='bottom'); ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[1].height = 58
for i, r in enumerate(rows, 2):
    v = {
      1: f'=RANK(G{i},$G$2:$G${last})',
      2: f'=IF(G{i}>=Parâmetros!$B$4,"A",IF(G{i}>=Parâmetros!$B$5,"B","C"))',
      3: r['n'], 4: r['uf'], 5: r['rg'], 6: r['arr'],
      7: f'=J{i}*S{i}', 8: f'=I{i}*Parâmetros!$B$6/100*S{i}',
      9: r['apt'], 10: r['aus'], 11: f'=J{i}/I{i}', 12: r['a22'], 13: f'=IF(K{i}>L{i},"sim","não")',
      14: r['dab'], 15: r['vl'], 16: r['vf'], 17: f'=O{i}/(O{i}+P{i})', 18: r['l2t'], 19: f'=2*Q{i}-1',
      20: r['i70'], 21: r['pp'], 22: r['pb'], 23: r['gov'], 24: r['ibge'], 25: r['pop'],
    }
    for j, val in v.items():
        c = ws.cell(row=i, column=j, value=val); c.number_format = cols[j - 1][2]
        c.font = blue if (j in (9, 10, 12, 14, 15, 16, 18, 20, 25) ) else body
ws.freeze_panes = 'D2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{last}'
ws['G1'].comment = Comment('Ausentes × vantagem de Lula. Votos líquidos se todos os ausentes votassem como os vizinhos. É um teto, não uma meta.', 'Método')
ws['H1'].comment = Comment('Eleitores aptos × meta de comparecimento (aba Parâmetros, padrão 1 p.p.) × vantagem de Lula. Uma meta realista.', 'Método')
ws['N1'].comment = Comment('Negativo = mais gente votou no 2º turno de 2022 do que no 1º. Sinal de que a cidade já respondeu a mobilização de 2º turno.', 'Método')

# ---------------- Parâmetros ----------------
wp = wb.create_sheet('Parâmetros')
wp['A1'] = 'Parâmetros (células amarelas podem ser editadas; a lista recalcula)'; wp['A1'].font = Font(name=F, bold=True, size=12)
for j, h in enumerate(['Parâmetro', 'Valor', 'Para que serve'], 1):
    c = wp.cell(row=3, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill
params = [('Prioridade A: saldo potencial mínimo (votos)', 2500, 'Cidades com saldo igual ou acima entram como A.'),
          ('Prioridade B: saldo potencial mínimo (votos)', 1000, 'Entre este valor e o da prioridade A, entram como B; abaixo, C.'),
          ('Meta de comparecimento a mais (p.p.)', 1, 'Usada na coluna "Saldo com +1 p.p.". Em 2022 a abstenção nacional caiu 0,36 p.p. entre os turnos.')]
for i, (a, b, c3) in enumerate(params, 4):
    wp.cell(row=i, column=1, value=a).font = body
    cb = wp.cell(row=i, column=2, value=b); cb.font = blue; cb.fill = yellow; cb.number_format = '#,##0' if b > 10 else '0.0'
    wp.cell(row=i, column=3, value=c3).font = body
wp.column_dimensions['A'].width = 46; wp.column_dimensions['B'].width = 10; wp.column_dimensions['C'].width = 90

# ---------------- Por estado ----------------
we = wb.create_sheet('Por estado')
hd = [('UF', 6, '@'), ('Região', 12, '@'), ('Cidades', 9, '#,##0'), ('Prioridade A', 11, '#,##0'), ('Eleitores aptos', 13, '#,##0'),
      ('Ausentes', 12, '#,##0'), ('Saldo potencial (votos)', 14, '#,##0'), ('Saldo com a meta de comparecimento (votos)', 16, '#,##0'), ('% do saldo da lista', 11, '0.0%')]
for j, (h, w, _) in enumerate(hd, 1):
    c = we.cell(row=1, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical='bottom')
    we.column_dimensions[get_column_letter(j)].width = w
we.row_dimensions[1].height = 44
ufs = {}
for r in rows: ufs.setdefault(r['uf'], [r['rg'], 0])[1] += r['saldo']
order = sorted(ufs, key=lambda u: -ufs[u][1])
tot = len(order) + 2
for i, u in enumerate(order, 2):
    vals = {1: u, 2: ufs[u][0], 3: f'=COUNTIF(Cidades!$D$2:$D${last},A{i})', 4: f'=COUNTIFS(Cidades!$D$2:$D${last},A{i},Cidades!$B$2:$B${last},"A")',
            5: f'=SUMIF(Cidades!$D$2:$D${last},A{i},Cidades!$I$2:$I${last})', 6: f'=SUMIF(Cidades!$D$2:$D${last},A{i},Cidades!$J$2:$J${last})',
            7: f'=SUMIF(Cidades!$D$2:$D${last},A{i},Cidades!$G$2:$G${last})', 8: f'=SUMIF(Cidades!$D$2:$D${last},A{i},Cidades!$H$2:$H${last})',
            9: f'=G{i}/$G${tot}'}
    for j, val in vals.items():
        c = we.cell(row=i, column=j, value=val); c.font = body; c.number_format = hd[j - 1][2]
we.cell(row=tot, column=1, value='Total').font = hdr_font
for j in range(3, 9):
    col = get_column_letter(j); c = we.cell(row=tot, column=j, value=f'=SUM({col}2:{col}{tot - 1})'); c.font = hdr_font; c.number_format = hd[j - 1][2]
c = we.cell(row=tot, column=9, value=f'=G{tot}/$G${tot}'); c.font = hdr_font; c.number_format = '0.0%'
we.freeze_panes = 'B2'

# ---------------- Como ler ----------------
wl = wb.create_sheet('Como ler')
lines = [
 ('Frente 2 · Mobilizar: cidades pró-Lula de 10 a 50 mil habitantes', 'title'),
 ('Quem entra', 'h'),
 (f'{N} municípios com 10.000 a 49.999 habitantes no Censo 2022 (IBGE) que são "área de Lula": Lula teve mais de 50% dos votos entre ele e Flávio no 1º turno de 2026 e venceu o 2º turno de 2022 no município.', ''),
 ('Ordem: do maior para o menor saldo potencial. A coluna Posição e a Prioridade se recalculam se você reordenar ou mudar os parâmetros.', ''),
 ('Colunas calculadas (fórmulas)', 'h'),
 ('Saldo potencial = ausentes × vantagem de Lula. Votos líquidos para Lula se todos os ausentes votassem como os vizinhos que votaram. É um teto.', ''),
 ('Saldo com +1 p.p. = eleitores aptos × meta (aba Parâmetros) × vantagem. Quanto rende, em votos líquidos, aumentar o comparecimento nessa medida.', ''),
 ('Vantagem de Lula = parcela de Lula − parcela de Flávio, só entre os dois. Um ausente numa cidade 80 a 20 rende 0,6 voto líquido; numa cidade 55 a 45, só 0,1.', ''),
 ('Abstenção subiu desde 2022? = sim quando a abstenção do 1º turno de 2026 passou a de 2022. Parte desses ausentes votou da última vez.', ''),
 ('Sinais de mobilização (dados de contexto)', 'h'),
 ('Variação da abstenção do 1º para o 2º turno de 2022: negativa quando mais gente votou no 2º turno. Indica que a cidade já respondeu a uma mobilização de 2º turno.', ''),
 ('Eleitores com 70 anos ou mais: para eles o voto é facultativo; costumam pesar na abstenção.', ''),
 ('Partido e campo do prefeito eleito em 2024 (esquerda, centro, direita): indica possíveis aliados locais.', ''),
 ('2º turno para governador no estado: mais um motivo para o eleitor ir votar em 25/10.', ''),
 ('Fontes', 'h'),
 ('TSE, Portal de Dados Abertos: votação por município e comparecimento, 1º turno de 2026 (arquivos de 08/10/2026); perfil do eleitorado 2026; candidatos eleitos em 2024.', ''),
 ('Resultados de 2022 (1º e 2º turno, abstenção): compilações conferidas com o TSE. População: Censo 2022 (IBGE, tabela 4709). Regiões geográficas imediatas: IBGE (2017).', ''),
 ('Números em azul são dados de entrada; em preto, fórmulas. Votos do exterior ficam de fora.', ''),
 ('Limites', 'h'),
 ('A conta supõe que os ausentes votariam como os vizinhos. Parte da abstenção é difícil de reverter (título numa cidade onde a pessoa já não mora, cadastro desatualizado).', ''),
 ('É uma lista de territórios, não de pessoas: diz onde concentrar a mobilização, não quem abordar.', ''),
 ('Uso', 'h'),
 ('Transporte de eleitores por candidatos, partidos ou apoiadores entre a véspera e o dia seguinte à eleição é proibido e é crime (Lei 6.091/1974, arts. 5º e 11). O transporte coletivo gratuito no dia da votação é obrigação do poder público (STF, ADPF 1.013; Resolução TSE 23.751/2026).', ''),
]
r = 1
for text, kind in lines:
    c = wl.cell(row=r, column=1, value=text)
    c.font = Font(name=F, bold=True, size=13) if kind == 'title' else Font(name=F, bold=True, size=11) if kind == 'h' else body
    c.alignment = Alignment(wrap_text=True, vertical='top')
    r += 2 if kind == 'title' else 1
    if kind == '' and lines.index((text, kind)) + 1 < len(lines) and lines[lines.index((text, kind)) + 1][1] == 'h': r += 1
wl.column_dimensions['A'].width = 120

wb.move_sheet('Como ler', offset=-3)
wb.active = 1
wb.save(OUT)
import shutil
shutil.copyfile(OUT, REPO / 'site/data' / OUT.name)
print(N, 'cidades; saldo', round(sum(r['saldo'] for r in rows)))
